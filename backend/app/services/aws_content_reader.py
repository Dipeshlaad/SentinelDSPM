import io
import json
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook


SUPPORTED_TEXT_TYPES = {
    ".txt",
    ".csv",
    ".json",
}

SUPPORTED_DOCUMENT_TYPES = {
    ".pdf",
    ".docx",
    ".xlsx",
}

def is_supported_object_type(object_key: str):
    extension = Path(object_key).suffix.lower()

    return (
        extension in SUPPORTED_TEXT_TYPES
        or extension in SUPPORTED_DOCUMENT_TYPES
    )

def read_s3_object_content(
    content: bytes,
    object_key: str,
):
    extension = Path(object_key).suffix.lower()

    if extension in SUPPORTED_TEXT_TYPES:
        return _read_text_content(content)

    if extension == ".pdf":
        return _read_pdf_content(content)

    if extension == ".docx":
        return _read_docx_content(content)

    if extension == ".xlsx":
        return _read_xlsx_content(content)

    return None


def _read_text_content(content: bytes):
    encodings = [
        "utf-8",
        "utf-16",
    ]

    for encoding in encodings:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue

    return None


def _read_pdf_content(content: bytes):
    text_parts = []

    with pymupdf.open(
        stream=content,
        filetype="pdf"
    ) as document:
        for page in document:
            text_parts.append(page.get_text())

    return "\n".join(text_parts)


def _read_docx_content(content: bytes):
    document = Document(
        io.BytesIO(content)
    )

    text_parts = []

    for paragraph in document.paragraphs:
        if paragraph.text:
            text_parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            row_values = []

            for cell in row.cells:
                if cell.text:
                    row_values.append(cell.text)

            if row_values:
                text_parts.append(" | ".join(row_values))

    return "\n".join(text_parts)


def _read_xlsx_content(content: bytes):
    workbook = load_workbook(
        filename=io.BytesIO(content),
        read_only=True,
        data_only=True,
    )

    text_parts = []

    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                values = []

                for value in row:
                    if value is not None:
                        values.append(str(value))

                if values:
                    text_parts.append(" | ".join(values))
    finally:
        workbook.close()

    return "\n".join(text_parts)